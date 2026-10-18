# -*- coding: utf-8 -*-
"""Build a REUDE Technologies–branded Word report from the two CompliQ intelligence docs."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(r"D:\temp\compliq-main")
DOCS = ROOT / "docs"
IMAGES = ROOT / "public" / "images"
OUT = DOCS / "REUDE_CompliQ_Competitive_Intelligence_Report.docx"

# Brand palette (sampled / established in CompliQ UI work)
ORANGE = RGBColor(0xFF, 0x6F, 0x00)
TEAL = RGBColor(0x0F, 0xA4, 0xAF)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_ROW = "F7F7F7"
HEADER_FILL = "FF6F00"
TEAL_FILL = "0FA4AF"


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, color=BLACK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade_cell(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    # Remove existing shd nodes so re-shading is clean
    for child in list(tcPr):
        if child.tag == qn("w:shd"):
            tcPr.remove(child)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        if edge in kwargs:
            element = OxmlElement(f"w:{edge}")
            element.set(qn("w:val"), kwargs[edge].get("val", "single"))
            element.set(qn("w:sz"), str(kwargs[edge].get("sz", 4)))
            element.set(qn("w:color"), kwargs[edge].get("color", "CCCCCC"))
            tcBorders.append(element)
    tcPr.append(tcBorders)


def add_horizontal_line(paragraph, color="FF6F00", size="18"):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr)
    run._r.append(fld_char_end)
    set_run_font(run, size=9, color=GRAY)


def paragraph(doc, text="", *, size=11, bold=False, italic=False, color=BLACK, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=8, space_before=0):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return p


def heading(doc, text, level=1):
    # Custom heading styles using brand colours
    sizes = {0: 22, 1: 16, 2: 13, 3: 12}
    colors = {0: ORANGE, 1: ORANGE, 2: TEAL, 3: BLACK}
    p = paragraph(
        doc,
        text,
        size=sizes.get(level, 12),
        bold=True,
        color=colors.get(level, BLACK),
        space_before=14 if level else 6,
        space_after=8,
    )
    if level in (0, 1):
        add_horizontal_line(p, color="FF6F00" if level <= 1 else "0FA4AF", size="12")
    return p


def bullets(doc, items, size=11):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        set_run_font(run, size=size, color=BLACK)


def numbered(doc, items, size=11):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        set_run_font(run, size=size, color=BLACK)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=9, bold=True, color=WHITE)
        shade_cell(cell, HEADER_FILL)
        set_cell_border(cell, top={"color": "FF6F00"}, bottom={"color": "FF6F00"}, left={"color": "FF6F00"}, right={"color": "FF6F00"})

    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(value))
            set_run_font(run, size=9, color=BLACK)
            if r_idx % 2 == 1:
                shade_cell(cell, LIGHT_ROW)

    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = w
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


def page_break(doc):
    doc.add_page_break()


def setup_header_footer(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.different_first_page_header_footer = True

    # First page header empty (cover)
    section.first_page_header.paragraphs[0].text = ""
    section.first_page_footer.paragraphs[0].text = ""

    header = section.header
    hp = header.paragraphs[0]
    hp.clear()
    logo = IMAGES / "reude R logo.png"
    if logo.exists():
        run = hp.add_run()
        run.add_picture(str(logo), width=Cm(1.0))
        hp.add_run("  ")
    run = hp.add_run("REUDE Technologies  |  CompliQ Competitive Intelligence")
    set_run_font(run, size=9, bold=True, color=ORANGE)
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_horizontal_line(hp, color="FF6F00", size="10")

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.clear()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("CONFIDENTIAL — Internal Use Only  ·  © REUDE Technologies  ·  Page ")
    set_run_font(run, size=8, color=GRAY)
    add_page_number(fp)


def cover_page(doc):
    # Top brand bar via table
    bar = doc.add_table(rows=1, cols=1)
    cell = bar.rows[0].cells[0]
    shade_cell(cell, "1A1A1A")
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo = IMAGES / "reude-logo.png"
    if logo.exists():
        run = p.add_run()
        run.add_picture(str(logo), width=Cm(8.5))
    else:
        run = p.add_run("REUDE TECHNOLOGIES")
        set_run_font(run, size=28, bold=True, italic=True, color=ORANGE)

    paragraph(doc, "", space_after=18)
    paragraph(doc, "COMPLIQ BY REUDE", size=12, bold=True, color=TEAL, align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(
        doc,
        "Competitive Intelligence & Product Benchmark Report",
        size=20,
        bold=True,
        color=BLACK,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=6,
    )
    paragraph(
        doc,
        "Wind Turbine Inspection Software — Market Gap Analysis\nand Nearthlab NearthWIND Pro REA Site & Video Study",
        size=12,
        color=GRAY,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=18,
    )

    line = paragraph(doc, "", align=WD_ALIGN_PARAGRAPH.CENTER)
    add_horizontal_line(line, color="FF6F00", size="18")

    # Document control box
    meta = [
        ("Document ID", "REUDE-COMPLIQ-CI-2026-10"),
        ("Classification", "Confidential — Internal"),
        ("Prepared for", "REUDE Technologies / CompliQ Product Planning"),
        ("Prepared by", "CompliQ Product Intelligence"),
        ("Issue date", date.today().strftime("%d %B %Y")),
        ("Version", "1.0"),
        ("Status", "Issued for review"),
        ("Parts", "A — Market Feature Gap Analysis\nB — Nearthlab NearthWIND Pro REA Report"),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(meta):
        c0, c1 = table.rows[i].cells
        c0.text = ""
        c1.text = ""
        r0 = c0.paragraphs[0].add_run(k)
        set_run_font(r0, size=10, bold=True, color=WHITE)
        shade_cell(c0, "1A1A1A")
        r1 = c1.paragraphs[0].add_run(v)
        set_run_font(r1, size=10, color=BLACK)
        shade_cell(c1, "FFF4E8" if i % 2 == 0 else "FFFFFF")
        c0.width = Cm(4.5)
        c1.width = Cm(11.5)

    paragraph(doc, "", space_after=24)
    paragraph(
        doc,
        "This document consolidates advertised competitor capabilities and REA-assisted evidence from public product sites. "
        "It is intended to guide CompliQ roadmap prioritisation. Competitor claims were not verified inside logged-in production systems unless stated.",
        size=9,
        italic=True,
        color=GRAY,
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    paragraph(doc, "www.reude.tech  ·  CompliQ by REUDE Technologies", size=9, color=ORANGE, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=20)
    page_break(doc)


def toc_page(doc):
    heading(doc, "Contents", level=0)
    items = [
        "Document control & purpose",
        "Part A — Wind Turbine Inspection Market Feature Gap Analysis",
        "    A1  Executive summary",
        "    A2  Sources reviewed",
        "    A3  Market standard",
        "    A4  Feature matrix vs CompliQ",
        "    A5  CompliQ strengths",
        "    A6  Priority build order",
        "    A7  Standards bar & 90-day outcomes",
        "Part B — Nearthlab NearthWIND Pro REA Site & Video Report",
        "    B1  Method and evidence",
        "    B2  Page architecture",
        "    B3  NearthWIND Pro findings",
        "    B4  Videos inventory",
        "    B5  Mobile & Zoomable context",
        "    B6  Technical interview synthesis",
        "    B7  Implications for CompliQ",
        "Closing recommendations",
    ]
    for item in items:
        paragraph(doc, item, size=11, color=BLACK if not item.startswith(" ") else GRAY, space_after=4)
    page_break(doc)


def part_a(doc):
    heading(doc, "Part A — Wind Turbine Inspection Market Feature Gap Analysis", level=0)
    paragraph(doc, "Document type: Competitive intelligence & product roadmap input", size=10, italic=True, color=GRAY)
    paragraph(doc, "Scope: Real-world blade / turbine inspection platforms versus CompliQ", size=10, italic=True, color=GRAY)
    paragraph(
        doc,
        "Method: Public product pages for market leaders, plus REA browser capture of key marketing sites; CompliQ assessed against the live dashboard in this project.",
        size=10,
        italic=True,
        color=GRAY,
        space_after=12,
    )

    heading(doc, "A1  Executive summary", level=1)
    paragraph(
        doc,
        "Market-leading wind inspection platforms do not sell “image folders.” They sell a closed operational loop:",
    )
    paragraph(
        doc,
        "Inspect → AI / expert confirm → locate on the blade → prioritise → plan repair → prove closure → retain history across campaigns.",
        bold=True,
        color=TEAL,
        space_after=10,
    )
    paragraph(doc, "CompliQ already has strong seeds for that loop:", bold=True)
    bullets(
        doc,
        [
            "Serial-first blade identity (e.g. Blade 5288)",
            "Interactive 3D turbine overview",
            "DNV-style CAT 1–5 severity language",
            "Open CV / manual ingest (software-first, not drone-locked)",
            "Fleet overview, asset detail, findings, and report surfaces",
        ],
    )
    paragraph(
        doc,
        "The main gaps versus market standard are lifecycle and evidence depth: repair work-order closure, true campaign-to-campaign progression, millimetre measurement on imagery, LPS / internal evidence on the same blade record, coverage gates, and audit-ready PDF reports linked to source frames.",
        space_before=8,
    )
    paragraph(
        doc,
        "Recommended strategy: Do not try to own drones, crawlers, or LPS sensors in the near term. Beat hardware-heavy vendors as the open system of record — serial blade passport, CAT SOP spine, vendor-agnostic ingest, repair closure, and defendable reports.",
        bold=True,
        space_before=6,
    )

    heading(doc, "A2  Sources reviewed", level=1)
    add_table(
        doc,
        ["Vendor / product", "What they sell", "Primary public source"],
        [
            ["SkySpecs Horizon (+ Foresight / SkyCrawler)", "Autonomous external inspection, internal crawler, blade management, repair planning, fleet analytics", "skyspecs.com/inspections"],
            ["Clobotics IRIS (+ IBIS / KIWI / Hummingbird)", "Unified external + internal + LPS evidence, AI screen → expert review, persistent blade record", "clobotics.com/iris"],
            ["TOPseven (INSPECTOR / NAVIGATOR / CLOUDseven / BEAT)", "Autonomous flight, AI severity, 3D twin, LPS analytics, fleet dashboard, audit PDF", "topseven.com"],
            ["SkyVisor Wind", "Autonomous DJI flights, AI defect detection, digital twin, FieldApp, tiered AI/expert plans", "skyvisor.ai"],
            ["Enervid", "Deep-zoom imagery, annotation, severity, repair assignment to closure, fleet map, PDF reports", "enervid.com/wind"],
            ["Aerones Visual Inspection Studio", "Internal crawler video + drone imagery, side-by-side defect tracking, AI anomaly assist", "aerones.com"],
            ["Perceptual Robotics (Dhalion)", "Autonomous capture, AI annotation, automated reporting, repair management, ISP portal", "perceptual-robotics.com"],
        ],
        col_widths=[Cm(5.2), Cm(7.0), Cm(4.0)],
    )
    paragraph(
        doc,
        "REA evidence (market scan): Browser scenario capture across TOPseven → Enervid → SkySpecs (evidence id ev_cae9d7ae…). Confirms live public pages; does not prove backend product behaviour.",
        size=9,
        italic=True,
        color=GRAY,
    )

    heading(doc, "A3  Market standard — what “good” looks like", level=1)
    numbered(
        doc,
        [
            "Asset-centric record — findings live on the turbine/blade, not in disconnected folders.",
            "Severity with action language — type + severity + recommended next step.",
            "Located damage — blade, surface/side, radius or station, preferably measurable size.",
            "Human-in-the-loop review — AI proposes; experts confirm.",
            "Comparable campaigns — new / stable / progressing over time.",
            "Repair path — assign, track, close with an audit trail.",
            "Fleet prioritisation — what to fix first across sites.",
            "Defendable outputs — reports linked to evidence, exportable, reviewable.",
            "Multi-evidence readiness — external RGB baseline; LPS / internal / thermal as add-ons.",
            "Access control — owners, ISPs, and repair crews see appropriate slices.",
        ],
    )

    heading(doc, "A4  Feature matrix — market vs CompliQ", level=1)
    add_table(
        doc,
        ["Feature", "Market standard", "Leaders", "CompliQ", "Gap"],
        [
            ["Finding → repair → close", "Assign, track to closure, audit trail", "Enervid, Horizon, IRIS", "Status/action fields; no full work-order board", "Partial"],
            ["Campaign progression", "Overlay inspections; new/stable/progressing", "TOPseven, Enervid, Clobotics, SkyVisor, Aerones", "Compare UI seeds only", "Partial"],
            ["Blade lifecycle record", "One blade identity across campaigns & vendors", "IRIS, Horizon", "Serial naming + findings; timeline incomplete", "Partial"],
            ["Measurable imagery", "Deep zoom + mm callouts", "Enervid, TOPseven", "Station photos / POC; limited calibrate measure", "Partial"],
            ["Photogrammetry / scan twin", "Turbine rebuilt from scan", "NAVIGATOR, SkyVisor, SkySpecs LiDAR", "Interactive generic 3D; not scan-derived", "Partial"],
            ["AI screen → expert confirm", "AI filters; expert confirms type/size/severity/location", "IRIS, TOPseven, SkyVisor Pro, Perceptual", "CV ingest exists; review queue incomplete", "Partial"],
            ["LPS evidence", "Pass/fail or located conductor fault on blade record", "BEAT, Hummingbird, SkyVisor Pro", "Not modelled", "Missing"],
            ["Internal + external evidence", "Crawler / internal with external imagery", "SkyCrawler, Aerones, KIWI", "External-oriented", "Missing"],
            ["Mission coverage verification", "Confirm all sides/stations before complete", "TOPseven, SkyVisor DroneApp", "Ingest without coverage gate", "Missing"],
            ["Audit-ready reports", "PDF tied to frames, charts, decisions", "TOPseven, Enervid, Perceptual", "Download exists; not fully evidence-generated", "Partial"],
            ["Fleet prioritisation analytics", "Frequency, severity, radius, urgency", "CLOUDseven, Clobotics, SkySpecs", "Metrics/charts; weak repair-risk ranking", "Partial"],
            ["Role-based access", "Owner / ISP / repair scoped views", "Enervid, SkyVisor, Perceptual", "Roles seeded; repair-vendor scope weak", "Partial"],
            ["SSO + open API / ingest", "Enterprise auth + any drone/CV pipeline", "SkyVisor Pro, Perceptual", "Open ingest strong; SSO/API incomplete", "Partial"],
        ],
    )

    heading(doc, "A5  What CompliQ already does well", level=1)
    numbered(
        doc,
        [
            "Serial blade identity — clearer for operators than letter-only A/B/C in many competitor demos.",
            "Interactive 3D overview — tangible asset context; competitors often lock twins to their capture stack.",
            "CAT 1–5 SOP language — aligns with serious O&M / engineering discussion of severity and action.",
            "Software-first / open ingest — can sit above any drone vendor, ISP, or edge AI.",
            "Inspection-flavoured product surface — fleet, asset, findings, report, ingest, analytics.",
        ],
    )

    heading(doc, "A6  Priority build order (to outpace the market)", level=1)
    priorities = [
        (
            "1. Defect lifecycle (finding → work order → verify → close)",
            "Every serious platform sells maintenance outcomes. Add a Repair Board: assignee, due date, cost, access window; require verified re-inspection before clearing CAT 4/5; auto-escalate if follow-up CAT is worse.",
        ),
        (
            "2. True campaign compare and progression",
            "Fingerprint defects by blade serial + surface + radius band. Status: New / Stable / Progressing / Resolved. Side-by-side station compare with delta severity.",
        ),
        (
            "3. Blade passport on the 3D model",
            "Click Blade 5288 → passport of findings, repairs, LPS/internal/thermal tabs, station photos. Pin defects at radius/surface on the existing 3D overview.",
        ),
        (
            "4. Evidence-linked audit reports",
            "Generate from live findings: CAT mix, annotated frames, action register, sign-off fields. Deep-link report rows back into the asset UI.",
        ),
        (
            "5. Multi-evidence without owning robots",
            "Evidence modules: External, Internal, Thermal, LPS, Rope-access. Partner ingest APIs; CompliQ owns decision + history.",
        ),
    ]
    for title, body in priorities:
        paragraph(doc, title, bold=True, color=TEAL, space_before=8, space_after=2)
        paragraph(doc, body, space_after=6)

    heading(doc, "A7  Standards bar & 90-day outcomes", level=1)
    paragraph(doc, "Standards bar to commit to", bold=True, color=ORANGE)
    bullets(
        doc,
        [
            "Condition & severity — CAT 1–5 as the single spine across fleet, asset, report, and repair boards.",
            "Evidence integrity — every finding linked to frame(s), blade serial, surface, radius/station, campaign, reviewer, final action.",
            "Comparable inspections — same stations/naming; coverage checklist per campaign.",
            "Insight to uptime — prioritise by lost-production risk and access window; clear CAT 4/5 only after verified repair evidence.",
        ],
    )
    paragraph(doc, "Suggested 90-day outcomes", bold=True, color=ORANGE, space_before=8)
    numbered(
        doc,
        [
            "Repair Board with assign / due / close / verify on existing defects.",
            "Campaign progression (New / Stable / Progressing) with side-by-side compare.",
            "Generated audit PDF from live findings + annotated evidence.",
        ],
    )
    paragraph(doc, "Supporting foundations in the same window:", bold=True, space_before=6)
    bullets(
        doc,
        [
            "Coverage checklist on ingest",
            "Blade passport panel from the 3D / component selector",
            "Explicit evidence-type field (external / thermal / internal / LPS)",
        ],
    )
    paragraph(doc, "What not to chase first", bold=True, color=ORANGE, space_before=8)
    add_table(
        doc,
        ["Tempting competitor feature", "Why delay"],
        [
            ["Building an autonomous drone stack", "Capital-heavy; SkySpecs / SkyVisor / TOPseven / Nearthlab already compete here"],
            ["Owning crawler / LPS hardware", "Partner or ingest instead; own the blade record"],
            ["Perfect photogrammetry twin day one", "Pin findings on current 3D + station photos first; add scan mesh later"],
            ["Broad multi-vertical expansion", "Finish wind blade lifecycle depth before spreading thin"],
        ],
    )
    page_break(doc)


def part_b(doc):
    heading(doc, "Part B — Nearthlab NearthWIND Pro REA Site & Video Report", level=0)
    paragraph(doc, "Primary URL: https://www.nearthlab.com/en/enterprise#nearthwind-pro", size=10, italic=True, color=GRAY)
    paragraph(doc, "Focus: NearthWIND Pro section, companion products on the same page, and embedded product videos", size=10, italic=True, color=GRAY, space_after=10)

    heading(doc, "B1  Method and evidence", level=1)
    add_table(
        doc,
        ["Target", "Method", "Result"],
        [
            ["Enterprise page (Pro anchor)", "REA capture_browser_scenario (Chromium)", "Complete DOM, accessibility tree, screenshots"],
            ["On-page product videos", "REA wait + timed playback on Nearthlab origin", "Three YouTube iframes discovered and exercised"],
            ["YouTube watch pages", "REA navigate + accessibility", "Titles, durations, channel, play state"],
            ["Direct /embed/ without referrer", "REA", "Failed — YouTube Error 153 (HTTP Referer required)"],
            ["Supporting pages", "Web fetch", "Homepage, Zoomable AI site, Pro2 page, industry interview"],
        ],
    )
    paragraph(doc, "Primary REA evidence IDs", bold=True, color=TEAL)
    add_table(
        doc,
        ["Capture", "Evidence ID (prefix)", "Notes"],
        [
            ["Enterprise Pro page", "ev_78705be2…", "Final URL with #nearthwind-pro"],
            ["Multi-video locator probe", "ev_d253e042…", "Revealed all three video titles/IDs"],
            ["Section walk Pro → Mobile → Zoomable", "ev_c3ca9c3e…", "Screenshots per section"],
            ["YouTube watch pages", "ev_b631287c…", "Metadata + playback controls"],
            ["Session export", "nearthlab_rea_evidence_bundle.json", "~13.7 MB, 7 records"],
        ],
    )
    paragraph(
        doc,
        "Confidence boundary — Observed: page structure, copy, specs, iframe URLs, embed behaviour, YouTube metadata, UI screenshots. "
        "Inferred: live Zoomable behaviour beyond marketing UI; absolute accuracy of mm/minute claims. "
        "Not verified: logged-in Zoomable, customer datasets, pricing, in-field flight performance.",
        size=9,
        italic=True,
        color=GRAY,
    )

    # Insert key screenshots if present
    shot_pro = DOCS / "nearth-section-play_pro.png"
    if shot_pro.exists():
        paragraph(doc, "Figure B-1. NearthWIND Pro hero section (REA screenshot)", size=9, italic=True, color=GRAY, space_before=8)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(shot_pro), width=Cm(15.5))

    heading(doc, "B2  Page architecture", level=1)
    paragraph(doc, "The Enterprise page is a three-product stack:")
    numbered(
        doc,
        [
            "NearthWIND Pro — purpose-built autonomous inspection drone hardware",
            "NearthWIND Mobile — app that automates small commercial drones",
            "Zoomable — cloud digital-twin / analytics SaaS",
        ],
    )
    paragraph(
        doc,
        "Accessibility tree confirms EN/KR language switch, left complementary nav with hash links (#nearthwind-pro, #nearthwind-mobile, #zoomable), and for each product: heading, tagline, description, still image, autoplaying muted looping YouTube iframe, then numbered feature sections. External Zoomable CTA: https://ai.zoomable.io",
        space_before=6,
    )

    heading(doc, "B3  NearthWIND Pro findings", level=1)
    paragraph(doc, "Positioning", bold=True, color=TEAL)
    bullets(
        doc,
        [
            "Eyebrow: Enterprise · Wind Turbine O&M",
            "Subtitle: Automated Wind Turbine Inspection Solution",
            "One-liner: Purpose-built autonomous drone hardware that automates wind turbine safety inspection",
        ],
    )
    paragraph(doc, "Capability pillars (on-page)", bold=True, color=TEAL, space_before=8)
    bullets(
        doc,
        [
            "01 Autonomous Flight — target recognition, autonomous path planning, collision avoidance, proximity standoff, AI-driven full-surface capture",
            "02 Precision — 45MP · 35mm full-frame camera for sharp, blur-free imagery",
            "03 Wind excellence — onshore/offshore; gusts up to 17 m/s (page claim); operable by unskilled personnel after brief training",
        ],
    )
    paragraph(doc, "Key specifications", bold=True, color=TEAL, space_before=8)
    add_table(
        doc,
        ["Specification", "Advertised value"],
        [
            ["Camera", "45MP · 35mm full-frame"],
            ["Wind resistance", "Gusts up to 17 m/s"],
            ["Operation", "Operable by unskilled personnel"],
        ],
    )
    paragraph(doc, "Reliability claims (marketing performance figures on-page)", bold=True, color=TEAL)
    bullets(
        doc,
        [
            "Beat baseline of 1 mm image quality / 25-minute inspection — claimed: 0.3 mm-class detection, 15-minute inspection, re-shoot rate under 1%",
            "600+ images per turbine at 45MP",
            "OEM customers named: Siemens Energy, Vestas, GE; 40+ countries; 7 years deployment",
            "Company narrative: 100,000+ blades inspected; ~6 hours rope access compressed to 15-minute autonomous flight",
        ],
    )

    heading(doc, "B4  Videos inventory (REA)", level=1)
    add_table(
        doc,
        ["Product", "iframe title", "YouTube ID", "Length / visibility"],
        [
            ["NearthWIND Pro", "NearthWIND Pro detail video", "DTCZW75EVRw", "~1:00 · Unlisted · title “드론뷰”"],
            ["NearthWIND Mobile", "NearthWIND Mobile detail video", "poN_UFjCXXw", "~0:45 · Unlisted · title “001 1”"],
            ["Zoomable", "Zoomable detail video", "GhAXEZsRqkQ", "~1:25 · Public · Nearthlab official channel"],
        ],
    )
    paragraph(doc, "Embed behaviour", bold=True, color=TEAL)
    bullets(
        doc,
        [
            "Lazy-loaded iframes with autoplay=1, mute=1, loop=1, controls=0 — silent looping product loops",
            "Standalone /embed/ without Nearthlab HTTP Referer returns YouTube Error 153",
            "Pro/Mobile loops are unlisted production assets; Zoomable is the official narrated product video",
            "Pro/Mobile accessibility trees expose no spoken script or chapter list — visual proof clips, not documentation",
        ],
    )

    shot_zoom = DOCS / "nearth-section-play_zoom.png"
    if shot_zoom.exists():
        paragraph(doc, "Figure B-2. Zoomable on-page UI (REA screenshot) — severity 1–5, blade/side filters, gallery, analytics", size=9, italic=True, color=GRAY, space_before=8)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(shot_zoom), width=Cm(15.5))

    heading(doc, "B5  NearthWIND Mobile & Zoomable (same page)", level=1)
    paragraph(doc, "NearthWIND Mobile", bold=True, color=TEAL)
    bullets(
        doc,
        [
            "Plug-and-play app: turns any small commercial drone into an automated inspection tool",
            "CES Innovation Awards 2022 badge on page visual",
            "Positioned for cost/logistics reduction and markets where dedicated industrial drones are hard to field",
        ],
    )
    paragraph(doc, "Zoomable (on-page + ai.zoomable.io)", bold=True, color=TEAL, space_before=8)
    bullets(
        doc,
        [
            "Seamless processing of hundreds of photos; report that once took >1 month now within a day (claim)",
            "Auto location + size of defects; millimetre zoom; year-over-year change tracking",
            "End-to-end SaaS: scheduling, analysis, time-series comparison, task assignment, automated reporting, 3D viewer, repair history",
            "UI observed: PRECISION DEFECT ANALYSIS & TRACKING — turbine T01, blades 1/2/3, sides LE/TE, Severity 1–5 chips, gallery, donut analytics",
            "Marketing site adds: multi-source ingest, stitching with gap/re-shoot guidance, damage tracking, API/modular SaaS",
        ],
    )

    heading(doc, "B6  Technical interview synthesis (company-stated)", level=1)
    paragraph(
        doc,
        "April 2025 Inside Unmanned Systems interview with Hajun Kwon (Head of Wind Business). Aligns with and expands the Enterprise page. Treat as company-stated, not independently measured.",
        size=9,
        italic=True,
        color=GRAY,
    )
    add_table(
        doc,
        ["Topic", "Company statement"],
        [
            ["Onboard vs cloud AI", "Onboard for autonomous flight; defect detection post-flight in Zoomable"],
            ["Navigation", "Vision-based GNC; not GPS waypoints; backup LiDAR; useful offshore / GPS-denied"],
            ["Severity model", "Proprietary Severity 1 (minor) → 5 (immediate shutdown), colour-coded in Zoomable"],
            ["Defect types", "Cracks, erosion, delamination + anomaly detection"],
            ["Pre-AI QA", "Blur/focus, lighting, coverage, metadata; drop redundant frames"],
            ["Predictive angle", "Damage progression; energy-loss estimate; repair timing guidance"],
            ["Evidence breadth", "External + LPS + internal inspections packaged together"],
            ["Offshore data logistics", "Images on SSD; upload from site office / logistics hub"],
            ["SCADA", "Zoomable can integrate as a module with customer SCADA"],
            ["Spec caution", "Page: 45MP / 17 m/s gusts; interview mentions 64MP / ~14 m/s gusts in places — numbers differ by source"],
        ],
    )

    heading(doc, "B7  Implications for CompliQ", level=1)
    paragraph(doc, "Where Nearthlab is ahead (relative to CompliQ’s current shape)", bold=True, color=ORANGE)
    numbered(
        doc,
        [
            "Closed capture→report loop with dedicated QA before defecting",
            "Digital twin UX tying 3D turbine to millimetre imagery and severity gallery",
            "Severity 1–5 colour system embedded in review UI",
            "Progression / energy-loss / repair timing analytics (claimed)",
            "Multi-evidence (external + LPS + internal) as one package",
            "Coverage / re-shoot discipline",
            "Task assignment + repair history inside Zoomable",
        ],
    )
    paragraph(doc, "Where CompliQ can still differentiate", bold=True, color=ORANGE, space_before=8)
    numbered(
        doc,
        [
            "Vendor-agnostic system of record — stay open to any drone / CV pipeline",
            "Serial blade passport (Blade 5288 style) as primary operator identity",
            "DNV/SOP CAT language as compliance spine versus proprietary severity taxonomy",
            "Software-first multi-ISP deployment without requiring a single airframe vendor",
            "Faster path to repair work-order closure and owner/ISP role separation if executed well",
        ],
    )
    paragraph(doc, "Concrete product lessons", bold=True, color=ORANGE, space_before=8)
    add_table(
        doc,
        ["Nearthlab pattern", "CompliQ opportunity"],
        [
            ["Severity chips 1–5 on gallery", "Keep CAT chips on findings gallery; make them filterable"],
            ["Blade + side (LE/TE) filters", "Strengthen side/surface filters in asset review"],
            ["3D twin with magnifier into photo", "Pin findings on CompliQ 3D overview → open station photo"],
            ["Auto report + repair history", "Evidence-linked report + Repair Board"],
            ["Coverage QA before AI", "Campaign coverage checklist in Ingest"],
            ["Mobile for frequent spot checks", "Support “spot campaign” ingest for low-cost revisit imagery"],
        ],
    )
    page_break(doc)


def closing(doc):
    heading(doc, "Closing recommendations", level=0)
    paragraph(
        doc,
        "Across both Part A (broad market) and Part B (Nearthlab deep dive), the commercial bar is consistent: a persistent, severity-ranked, evidence-linked blade record that drives repair and proves closure.",
    )
    paragraph(doc, "Immediate CompliQ focus", bold=True, color=ORANGE, space_before=8)
    numbered(
        doc,
        [
            "Ship Repair Board (assign → verify → close) on existing defects.",
            "Ship campaign progression with side-by-side compare.",
            "Ship evidence-linked audit PDF generation.",
            "Add coverage checklist + evidence-type fields + blade passport on the 3D model.",
            "Treat LPS / internal / thermal as ingestable evidence modules — do not build drones first.",
        ],
    )
    paragraph(
        doc,
        "Nearthlab’s stack (Pro + Mobile + Zoomable) is the clearest single-vendor illustration of the end-to-end loop CompliQ should match in software while remaining capture-vendor-agnostic.",
        space_before=10,
    )

    line = paragraph(doc, "", space_before=18)
    add_horizontal_line(line, color="FF6F00", size="18")
    paragraph(doc, "— End of document —", align=WD_ALIGN_PARAGRAPH.CENTER, color=GRAY, size=10, space_before=8)
    paragraph(doc, "REUDE Technologies  ·  CompliQ  ·  Confidential", align=WD_ALIGN_PARAGRAPH.CENTER, color=ORANGE, size=9, bold=True)


def main():
    doc = Document()
    # Default style
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = BLACK

    setup_header_footer(doc)
    cover_page(doc)
    toc_page(doc)

    heading(doc, "Document purpose", level=1)
    paragraph(
        doc,
        "This report combines two intelligence workstreams produced for CompliQ product planning: (A) a cross-market feature gap analysis against leading wind inspection platforms, and (B) a REA-assisted deep dive of Nearthlab’s Enterprise / NearthWIND Pro public site and embedded videos. Together they define the market standard CompliQ should meet and a concrete near-term build order.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    page_break(doc)

    part_a(doc)
    part_b(doc)
    closing(doc)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Wrote {OUT}")
    print(f"Size: {OUT.stat().st_size} bytes")


if __name__ == "__main__":
    main()
